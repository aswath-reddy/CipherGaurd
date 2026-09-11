"""
CipherGuard Project Review Presentation Generator
Generates a professional 16:9 widescreen presentation deck covering:
- Software Requirements Specification (SRS)
- System Design
- Class Diagram
- Sequence Diagram
- State Chart Diagram
- Deployment Diagram
- Traceability Matrix & Review Summary
"""

import os
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = "docs/generated_diagrams"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# -------------------------------------------------------------
# Color Palette
# -------------------------------------------------------------
NAVY = "#0F172A"       # Slate 900
DARK_BLUE = "#1E3A8A"  # Blue 900
ACCENT_BLUE = "#2563EB"# Blue 600
CYAN = "#06B6D4"       # Cyan 500
TEAL = "#0D9488"       # Teal 600
BG_LIGHT = "#F8FAFC"   # Slate 50
CARD_BG = "#FFFFFF"    # White
BORDER_COLOR = "#CBD5E1"# Slate 300
TEXT_DARK = "#0F172A"
TEXT_MUTED = "#475569"
GREEN = "#16A34A"
AMBER = "#D97706"
RED = "#DC2626"


# -------------------------------------------------------------
# 1. Matplotlib Diagram Generators
# -------------------------------------------------------------

def create_system_design_diagram():
    fig, ax = plt.subplots(figsize=(10, 5.5), dpi=300)
    ax.set_facecolor(BG_LIGHT)
    fig.patch.set_facecolor(BG_LIGHT)
    ax.axis('off')

    # Draw Subsystem Boxes
    boxes = [
        ("1. Ingestion Layer", ["DirectInput (User prompts)", "IndirectInput (Docs, Tools, Web)", "Surface Normalizer & Provenance"], 0.05, 0.55, 0.26, 0.38, DARK_BLUE),
        ("2. Classification Layer", ["Model Family A (TF-IDF + LogReg)", "Model Family B (MiniLM + MLP)", "DeBERTa-v3 Zero-Shot Model", "Calibrated Score Fusion"], 0.37, 0.55, 0.26, 0.38, ACCENT_BLUE),
        ("3. Policy Engine", ["Hot-Reloadable policy.json", "Multi-Surface Threshold Matching", "Decision Resolution (BLOCK/REVIEW/ALLOW)"], 0.69, 0.55, 0.26, 0.38, TEAL),
        ("4. Attribution Engine", ["Iterative Beam Search (k=5, L=6)", "Minimal Token Set Identification", "Sanitized Safe Text Generation"], 0.21, 0.08, 0.26, 0.38, AMBER),
        ("5. Persistence & Interface", ["SQLite WAL DB (cipherguard_audit.db)", "FastAPI REST Server (:8000)", "Streamlit Dashboard (:8501)"], 0.53, 0.08, 0.26, 0.38, GREEN),
    ]

    for title, items, x, y, w, h, col in boxes:
        # Header card
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.03",
                                      facecolor=CARD_BG, edgecolor=col, linewidth=2, transform=ax.transAxes)
        ax.add_patch(rect)
        ax.text(x + 0.02, y + h - 0.05, title, fontsize=11, fontweight='bold', color=col, transform=ax.transAxes)
        
        # Items
        item_y = y + h - 0.11
        for it in items:
            ax.text(x + 0.03, item_y, f"• {it}", fontsize=9, color=TEXT_DARK, transform=ax.transAxes)
            item_y -= 0.065

    # Connectors
    arrow_props = dict(arrowstyle="->", color=NAVY, lw=2)
    ax.annotate("", xy=(0.37, 0.74), xytext=(0.31, 0.74), arrowprops=arrow_props, xycoords="axes fraction")
    ax.annotate("", xy=(0.69, 0.74), xytext=(0.63, 0.74), arrowprops=arrow_props, xycoords="axes fraction")
    ax.annotate("", xy=(0.34, 0.46), xytext=(0.78, 0.55), arrowprops=dict(arrowstyle="->", color=RED, lw=1.8, linestyle="dashed"), xycoords="axes fraction")
    ax.text(0.57, 0.49, "If BLOCK / REVIEW", fontsize=8, color=RED, fontweight='bold', transform=ax.transAxes)
    ax.annotate("", xy=(0.53, 0.27), xytext=(0.47, 0.27), arrowprops=arrow_props, xycoords="axes fraction")

    plt.tight_layout()
    path = os.path.join(OUTPUT_DIR, "system_design.png")
    plt.savefig(path, bbox_inches='tight', facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path


def create_class_diagram():
    fig, ax = plt.subplots(figsize=(11, 6), dpi=300)
    ax.set_facecolor(BG_LIGHT)
    fig.patch.set_facecolor(BG_LIGHT)
    ax.axis('off')

    classes = [
        ("BaseClassifier (Abstract)", ["+predict_proba(texts)*"], 0.04, 0.68, 0.27, 0.24, DARK_BLUE),
        ("ModelFamilyA", ["-vectorizer: Tfidf", "-classifier: LogReg", "+predict_proba()", "+train()"], 0.04, 0.36, 0.27, 0.26, ACCENT_BLUE),
        ("RiskAggregator", ["-models: List[BaseClassifier]", "-weights: Dict[str, float]", "+aggregate(text, surface)"], 0.36, 0.68, 0.30, 0.26, TEAL),
        ("DecisionEngine", ["-loader: PolicyLoader", "+evaluate(scores, surface)"], 0.70, 0.68, 0.26, 0.24, TEAL),
        ("PolicyLoader", ["-policy_path: str", "-cached_config: PolicyConfig", "+load_policy()", "+reload()"], 0.70, 0.36, 0.26, 0.26, NAVY),
        ("CipherGuardRouter", ["-aggregator: RiskAggregator", "-decision_engine: DecisionEngine", "-audit_logger: AuditLogger", "+route(input_data) RouterOutput"], 0.36, 0.28, 0.30, 0.32, DARK_BLUE),
        ("AttributionEngine", ["+run_removal_beam_search()", "+format_contrastive_explanation()"], 0.04, 0.04, 0.27, 0.24, AMBER),
        ("AuditLogger", ["-db_path: str", "+log_decision()", "+get_recent_logs()"], 0.70, 0.04, 0.26, 0.24, GREEN),
    ]

    for title, methods, x, y, w, h, col in classes:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.015,rounding_size=0.02",
                                      facecolor=CARD_BG, edgecolor=col, linewidth=1.8, transform=ax.transAxes)
        ax.add_patch(rect)
        # Class Header
        header = patches.Rectangle((x, y + h - 0.06), w, 0.06, facecolor=col, transform=ax.transAxes)
        ax.add_patch(header)
        ax.text(x + w/2, y + h - 0.04, title, fontsize=9.5, fontweight='bold', color='white', ha='center', transform=ax.transAxes)
        
        my = y + h - 0.10
        for m in methods:
            ax.text(x + 0.015, my, m, fontsize=8, color=TEXT_DARK, transform=ax.transAxes)
            my -= 0.045

    # Relationships
    ax.annotate("", xy=(0.17, 0.68), xytext=(0.17, 0.62), arrowprops=dict(arrowstyle="->", color=DARK_BLUE, lw=1.5), xycoords="axes fraction")
    ax.text(0.18, 0.64, "inherits", fontsize=8, color=TEXT_MUTED, transform=ax.transAxes)
    
    ax.annotate("", xy=(0.36, 0.78), xytext=(0.31, 0.78), arrowprops=dict(arrowstyle="->", color=TEAL, lw=1.5), xycoords="axes fraction")
    ax.text(0.32, 0.80, "aggregates", fontsize=8, color=TEXT_MUTED, transform=ax.transAxes)

    ax.annotate("", xy=(0.70, 0.78), xytext=(0.66, 0.78), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.5), xycoords="axes fraction")
    
    ax.annotate("", xy=(0.51, 0.60), xytext=(0.51, 0.68), arrowprops=dict(arrowstyle="->", color=DARK_BLUE, lw=1.5), xycoords="axes fraction")
    ax.annotate("", xy=(0.83, 0.68), xytext=(0.83, 0.62), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.5), xycoords="axes fraction")

    plt.tight_layout()
    path = os.path.join(OUTPUT_DIR, "class_diagram.png")
    plt.savefig(path, bbox_inches='tight', facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path


def create_sequence_diagram():
    fig, ax = plt.subplots(figsize=(11, 6), dpi=300)
    ax.set_facecolor(BG_LIGHT)
    fig.patch.set_facecolor(BG_LIGHT)
    ax.axis('off')

    lifelines = [
        ("Client", 0.08),
        ("Router", 0.25),
        ("Classifier", 0.44),
        ("PolicyEngine", 0.62),
        ("BeamSearch", 0.80),
        ("AuditDB", 0.94),
    ]

    # Draw lifelines
    for name, x in lifelines:
        box = patches.FancyBboxPatch((x - 0.06, 0.88), 0.12, 0.08, boxstyle="round,pad=0.01", facecolor=DARK_BLUE, edgecolor='none')
        ax.add_patch(box)
        ax.text(x, 0.92, name, fontsize=9.5, fontweight='bold', color='white', ha='center', va='center')
        ax.axvline(x=x, ymin=0.08, ymax=0.88, color=BORDER_COLOR, linestyle="--", lw=1.5)

    messages = [
        (0.08, 0.25, 0.82, "1. route(prompt, surface='direct')", ACCENT_BLUE, False),
        (0.25, 0.44, 0.73, "2. aggregate(text, surface)", TEAL, False),
        (0.44, 0.25, 0.65, "3. return risk_vector [0.88]", TEAL, True),
        (0.25, 0.62, 0.57, "4. evaluate(risk_vector)", DARK_BLUE, False),
        (0.62, 0.25, 0.49, "5. DecisionResult (BLOCK, rule_id)", DARK_BLUE, True),
        (0.25, 0.80, 0.40, "6. run_removal_beam_search(text, thresh)", AMBER, False),
        (0.80, 0.25, 0.32, "7. BeamResult (removed_tokens, is_flipped)", AMBER, True),
        (0.25, 0.94, 0.24, "8. log_decision(audit_record)", GREEN, False),
        (0.94, 0.25, 0.18, "9. audit_id=104", GREEN, True),
        (0.25, 0.08, 0.11, "10. RouterOutput (BLOCK, sanitized_text)", RED, True),
    ]

    for x1, x2, y, label, col, is_ret in messages:
        ls = "--" if is_ret else "-"
        ax.annotate("", xy=(x2, y), xytext=(x1, y),
                    arrowprops=dict(arrowstyle="->", color=col, lw=1.6, linestyle=ls))
        mid_x = (x1 + x2) / 2
        ax.text(mid_x, y + 0.02, label, fontsize=8, color=TEXT_DARK, ha='center', fontweight='semibold')

    plt.tight_layout()
    path = os.path.join(OUTPUT_DIR, "sequence_diagram.png")
    plt.savefig(path, bbox_inches='tight', facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path


def create_statechart_diagram():
    fig, ax = plt.subplots(figsize=(11, 5.5), dpi=300)
    ax.set_facecolor(BG_LIGHT)
    fig.patch.set_facecolor(BG_LIGHT)
    ax.axis('off')

    states = [
        ("[*] Start", 0.03, 0.45, 0.09, 0.12, NAVY),
        ("Ingestion & Normalization", 0.16, 0.42, 0.18, 0.18, DARK_BLUE),
        ("Classification Ensemble", 0.38, 0.42, 0.18, 0.18, ACCENT_BLUE),
        ("Policy Rule Resolution", 0.60, 0.42, 0.18, 0.18, TEAL),
        ("ALLOW State\n(Forward to LLM)", 0.83, 0.68, 0.15, 0.22, GREEN),
        ("Flagged State (BLOCK/REVIEW)\nRemoval Beam Search", 0.81, 0.18, 0.17, 0.25, AMBER),
        ("Audit Logging & Return\n(SQLite Commit)", 0.50, 0.04, 0.20, 0.18, DARK_BLUE),
    ]

    for title, x, y, w, h, col in states:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.015,rounding_size=0.03",
                                      facecolor=CARD_BG, edgecolor=col, linewidth=2, transform=ax.transAxes)
        ax.add_patch(rect)
        ax.text(x + w/2, y + h/2, title, fontsize=9, fontweight='bold', color=col, ha='center', va='center', transform=ax.transAxes)

    # State Transitions
    ax.annotate("", xy=(0.16, 0.51), xytext=(0.12, 0.51), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.8), xycoords="axes fraction")
    ax.annotate("", xy=(0.38, 0.51), xytext=(0.34, 0.51), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.8), xycoords="axes fraction")
    ax.annotate("", xy=(0.60, 0.51), xytext=(0.56, 0.51), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.8), xycoords="axes fraction")
    
    # Allow branch
    ax.annotate("", xy=(0.83, 0.77), xytext=(0.74, 0.60), arrowprops=dict(arrowstyle="->", color=GREEN, lw=1.8), xycoords="axes fraction")
    ax.text(0.74, 0.70, "Action == ALLOW", fontsize=8, color=GREEN, fontweight='bold', transform=ax.transAxes)

    # Block branch
    ax.annotate("", xy=(0.81, 0.32), xytext=(0.74, 0.42), arrowprops=dict(arrowstyle="->", color=RED, lw=1.8), xycoords="axes fraction")
    ax.text(0.73, 0.35, "Action == BLOCK/REVIEW", fontsize=8, color=RED, fontweight='bold', transform=ax.transAxes)

    # Converge to Logging
    ax.annotate("", xy=(0.66, 0.22), xytext=(0.83, 0.68), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.5, linestyle="dotted"), xycoords="axes fraction")
    ax.annotate("", xy=(0.70, 0.13), xytext=(0.81, 0.22), arrowprops=dict(arrowstyle="->", color=NAVY, lw=1.5, linestyle="dotted"), xycoords="axes fraction")

    plt.tight_layout()
    path = os.path.join(OUTPUT_DIR, "statechart_diagram.png")
    plt.savefig(path, bbox_inches='tight', facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path


def create_deployment_diagram():
    fig, ax = plt.subplots(figsize=(11, 5.8), dpi=300)
    ax.set_facecolor(BG_LIGHT)
    fig.patch.set_facecolor(BG_LIGHT)
    ax.axis('off')

    # Subgraph Tiers
    tiers = [
        ("Client Tier", ["Admin Browser (:8501)", "Agent Application (SDK)"], 0.03, 0.60, 0.22, 0.32, DARK_BLUE),
        ("Ingress / Proxy Tier", ["Reverse Proxy (Nginx / TLS)", "Port 8000 / 8501 Dispatch"], 0.29, 0.60, 0.22, 0.32, ACCENT_BLUE),
        ("Application & Policy Tier", ["Uvicorn ASGI (:8000)", "FastAPI Gateway (api/main.py)", "CipherGuardRouter Core", "DecisionEngine & BeamSearch"], 0.55, 0.48, 0.42, 0.44, TEAL),
        ("Model Worker Tier (GPU/CPU)", ["ModelFamilyA (TF-IDF)", "ModelFamilyB (MiniLM)", "DeBERTa-v3 Worker"], 0.03, 0.10, 0.48, 0.36, AMBER),
        ("Persistence Tier", ["policy.json (Hot-Reload)", "cipherguard_audit.db (SQLite WAL)"], 0.55, 0.10, 0.42, 0.32, GREEN),
    ]

    for title, items, x, y, w, h, col in tiers:
        rect = patches.FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.015,rounding_size=0.03",
                                      facecolor=CARD_BG, edgecolor=col, linewidth=2, transform=ax.transAxes)
        ax.add_patch(rect)
        ax.text(x + 0.02, y + h - 0.06, title, fontsize=10.5, fontweight='bold', color=col, transform=ax.transAxes)
        
        iy = y + h - 0.12
        for it in items:
            ax.text(x + 0.03, iy, f"• {it}", fontsize=9, color=TEXT_DARK, transform=ax.transAxes)
            iy -= 0.06

    # Connections
    ax.annotate("", xy=(0.29, 0.76), xytext=(0.25, 0.76), arrowprops=dict(arrowstyle="->", color=NAVY, lw=2), xycoords="axes fraction")
    ax.annotate("", xy=(0.55, 0.76), xytext=(0.51, 0.76), arrowprops=dict(arrowstyle="->", color=NAVY, lw=2), xycoords="axes fraction")
    ax.annotate("", xy=(0.55, 0.26), xytext=(0.51, 0.26), arrowprops=dict(arrowstyle="<->", color=GREEN, lw=1.8), xycoords="axes fraction")
    ax.annotate("", xy=(0.60, 0.48), xytext=(0.40, 0.46), arrowprops=dict(arrowstyle="->", color=AMBER, lw=1.8), xycoords="axes fraction")
    ax.text(0.43, 0.50, "Batch Inference", fontsize=8, color=AMBER, fontweight='bold', transform=ax.transAxes)

    plt.tight_layout()
    path = os.path.join(OUTPUT_DIR, "deployment_diagram.png")
    plt.savefig(path, bbox_inches='tight', facecolor=fig.get_facecolor(), edgecolor='none')
    plt.close()
    return path


# -------------------------------------------------------------
# 2. PowerPoint Slide Deck Builder
# -------------------------------------------------------------

def build_presentation(diagram_paths):
    prs = Presentation()
    # Set 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="CIPHERGUARD • PROJECT REVIEW 2"):
        # Header banner shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate 900
        banner.line.fill.background()

        # Category pill text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = RGBColor(6, 182, 212) # Cyan

        # Title text
        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.5), Inches(0.65))
        tf2 = tx_box2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(255, 255, 255)

    def add_card(slide, left, top, width, height, bg_color=RGBColor(255, 255, 255), border_color=RGBColor(203, 213, 225)):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ---------------------------------------------------------
    # SLIDE 1: Title Slide
    # ---------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg1.line.fill.background()

    # Decorative Cyan Line
    line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.5), Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(6, 182, 212)
    line.line.fill.background()

    # Title & Subtitle box
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(2.5))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "CipherGuard"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = tf.add_paragraph()
    p_sub.text = "A Policy-Aware Safety Router with Removal-Based Token Attribution for LLM Applications"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(147, 197, 253) # Light blue

    p_desc = tf.add_paragraph()
    p_desc.text = "Project Review 2: Software Requirements Specification (SRS) & System Design Specification"
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = RGBColor(148, 163, 184) # Slate 400

    # Meta card
    add_card(slide1, Inches(1.2), Inches(5.0), Inches(10.9), Inches(1.8), RGBColor(30, 41, 59), RGBColor(51, 65, 85))
    mbox = slide1.shapes.add_textbox(Inches(1.5), Inches(5.15), Inches(10.3), Inches(1.5))
    mtf = mbox.text_frame
    
    mp1 = mtf.paragraphs[0]
    mp1.text = "Keshav Memorial Institute of Technology (KMIT), Hyderabad • Department of CSE"
    mp1.font.size = Pt(12)
    mp1.font.bold = True
    mp1.font.color.rgb = RGBColor(6, 182, 212)

    mp2 = mtf.add_paragraph()
    mp2.text = "Under the Guidance of: B. Shailesh"
    mp2.font.size = Pt(13)
    mp2.font.color.rgb = RGBColor(255, 255, 255)

    mp3 = mtf.add_paragraph()
    mp3.text = "Team Members: Shreya Namdeo • Shashank Reddy Yasa • Adi Aswatha Reddy • B. Abhinav"
    mp3.font.size = Pt(13)
    mp3.font.bold = True
    mp3.font.color.rgb = RGBColor(226, 232, 240)

    # ---------------------------------------------------------
    # SLIDE 2: Review Agenda & Scope
    # ---------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Project Review 2: Scope & Agenda")

    agenda_items = [
        ("1. Software Requirements Specification (SRS)", "Functional requirements (FR-1 to FR-5), non-functional constraints, and IEEE 830 compliance.", ACCENT_BLUE),
        ("2. System Design & Architectural Patterns", "Decoupled gateway pipeline, threat models, and sub-50ms inference design.", TEAL),
        ("3. Class Diagram & OOP Structure", "Component hierarchy, polymorphism, and dynamic policy loader relationship.", DARK_BLUE),
        ("4. Sequence Diagrams", "End-to-end request routing, beam search attribution, and runtime policy hot-reloading.", AMBER),
        ("5. State Chart Diagram", "Lifecycle state machine from input ingestion to fail-closed safety resolution.", GREEN),
        ("6. Deployment Diagram & Physical Topology", "Process boundaries, ports (:8000, :8501), inference workers, and SQLite storage.", NAVY),
    ]

    for idx, (title, desc, col) in enumerate(agenda_items):
        row = idx // 2
        cidx = idx % 2
        left = Inches(0.8 + cidx * 5.9)
        top = Inches(1.5 + row * 1.8)

        add_card(slide2, left, top, Inches(5.6), Inches(1.55), RGBColor(255, 255, 255), RGBColor(226, 232, 240))
        
        # Pill on card
        pill = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.2), top + Inches(0.2), Inches(0.08), Inches(1.15))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(13, 148, 136) if col == TEAL else RGBColor(37, 99, 235)
        pill.line.fill.background()

        tbox = slide2.shapes.add_textbox(left + Inches(0.4), top + Inches(0.15), Inches(5.0), Inches(1.2))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(15, 23, 42)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 3: SRS - Functional Requirements
    # ---------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Software Requirements: Functional Modules (FR-1 to FR-5)")

    fr_cards = [
        ("FR-1: Dual-Surface Ingestion", [
            "Accepts direct user prompts (surface='direct')",
            "Ingests indirect context: docs, tools, emails (surface='indirect')",
            "Preserves provenance metadata across the routing lifecycle"
        ]),
        ("FR-2: Risk Scoring & Ensembles", [
            "Ensembles Model Family A (TF-IDF+LogReg) & Family B (MLP)",
            "Optional zero-shot transformer scoring via DeBERTa-v3",
            "Calibrated risk outputs: Jailbreak, Injection, PII, Tools, Hate"
        ]),
        ("FR-3: Dynamic Policy Engine", [
            "JSON-configurable policy matrix (policy.json)",
            "Runtime threshold hot-reloading with zero retraining",
            "Hierarchical rule matching (surface-specific -> wildcard *)"
        ]),
        ("FR-4: Removal-Based Attribution", [
            "Iterative beam search (k=5, L=6) on BLOCK or REVIEW",
            "Identifies minimal token subset flipping decision to PASS",
            "Computes delta score and generates sanitized safe prompt"
        ]),
        ("FR-5: Audit Trail & Interfaces", [
            "Persistent SQLite audit logging with WAL integrity",
            "FastAPI REST endpoints (/route, /policy, /logs, /explain)",
            "Interactive Streamlit web dashboard for administration"
        ]),
    ]

    for idx, (title, bullets) in enumerate(fr_cards):
        if idx < 3:
            left = Inches(0.8 + idx * 3.9)
            top = Inches(1.5)
            w = Inches(3.7)
            h = Inches(2.6)
        else:
            left = Inches(1.8 + (idx - 3) * 5.0)
            top = Inches(4.35)
            w = Inches(4.7)
            h = Inches(2.6)

        add_card(slide3, left, top, w, h)
        tbox = slide3.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), w - Inches(0.4), h - Inches(0.3))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 58, 138)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.size = Pt(10)
            pb.font.color.rgb = RGBColor(51, 65, 85)

    # ---------------------------------------------------------
    # SLIDE 4: SRS - Non-Functional Requirements & Constraints
    # ---------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Software Requirements: Non-Functional Requirements (NFR)")

    nfr_items = [
        ("NFR-1: Performance & Low Latency", [
            "Lightweight classifier routing must complete in <= 50ms.",
            "Removal beam search attribution completes in <= 400ms for typical inputs (<= 60 tokens).",
            "Zero latency overhead added to benign prompts (attribution runs strictly on flagged inputs)."
        ]),
        ("NFR-2: Security & Fail-Closed Integrity", [
            "Strict Pydantic schema validation on all dynamic policy updates.",
            "Parameterized SQLite statements to mitigate internal injection risks.",
            "Fail-closed architecture: unhandled classifier exceptions default to BLOCK."
        ]),
        ("NFR-3: Reliability & High Availability", [
            "Atomic hot-reloading ensures 99.9% uptime during runtime policy edits.",
            "File-watcher fallback protects system from corrupted or malformed policy JSON files."
        ]),
        ("NFR-4: Auditability & Explainability", [
            "Every routing transaction logs inputs, risk vectors, rule evaluations, and latency.",
            "Contrastive explanations expose minimal trigger tokens and delta risk reduction."
        ]),
    ]

    for idx, (title, bullets) in enumerate(nfr_items):
        r = idx // 2
        c = idx % 2
        left = Inches(0.8 + c * 5.9)
        top = Inches(1.5 + r * 2.7)

        add_card(slide4, left, top, Inches(5.6), Inches(2.45))
        tbox = slide4.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), Inches(5.1), Inches(2.0))
        tf = tbox.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(13, 148, 136)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.size = Pt(10.5)
            pb.font.color.rgb = RGBColor(51, 65, 85)

    # ---------------------------------------------------------
    # SLIDE 5: System Design & Subsystem Architecture
    # ---------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "System Design: Layered Architecture & Subsystems")

    # Insert Generated Diagram Image on Left
    slide5.shapes.add_picture(diagram_paths["system_design"], Inches(0.8), Inches(1.5), width=Inches(7.4))

    # Explanatory Card on Right
    add_card(slide5, Inches(8.5), Inches(1.5), Inches(4.0), Inches(5.4))
    tbox = slide5.shapes.add_textbox(Inches(8.7), Inches(1.7), Inches(3.6), Inches(5.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Key Architectural Highlights"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    points = [
        ("Gateway Pattern", "Intercepts direct & indirect inputs before generative LLM execution."),
        ("Score Fusion", "Ensembles lexical (TF-IDF), dense neural (MLP), and transformer (DeBERTa-v3) classifiers."),
        ("Decoupled Policy", "Decouples threshold policies from model weights—no retraining required for policy changes."),
        ("Conditional Search", "Beam search attribution activates exclusively on BLOCK/REVIEW decisions."),
        ("Audit Persistence", "All decisions and contrastive rationale are committed to SQLite WAL storage.")
    ]
    for h, d in points:
        ph = tf.add_paragraph()
        ph.text = f"• {h}:"
        ph.font.size = Pt(11)
        ph.font.bold = True
        ph.font.color.rgb = RGBColor(37, 99, 235)
        pd = tf.add_paragraph()
        pd.text = f"   {d}"
        pd.font.size = Pt(10)
        pd.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 6: Class Diagram
    # ---------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "System Design: Object-Oriented Class Diagram")

    slide6.shapes.add_picture(diagram_paths["class_diagram"], Inches(0.8), Inches(1.4), width=Inches(7.8))

    add_card(slide6, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.5))
    tbox = slide6.shapes.add_textbox(Inches(9.0), Inches(1.6), Inches(3.3), Inches(5.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Class Design Notes"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    class_notes = [
        ("Polymorphic Classifiers", "BaseClassifier enforces predict_proba() contract across ModelFamilyA, ModelFamilyB, and DebertaClassifier."),
        ("Risk Aggregation", "RiskAggregator fuses model predictions using calibrated weighting."),
        ("Policy & Decision", "DecisionEngine evaluates fused scores against PolicyLoader cached rules."),
        ("Core Orchestrator", "CipherGuardRouter orchestrates ingestion, classification, policy check, attribution, and audit logging."),
        ("Audit Trail", "AuditLogger persists decision records and sanitized texts to SQLite.")
    ]
    for h, d in class_notes:
        ph = tf.add_paragraph()
        ph.text = f"• {h}:"
        ph.font.size = Pt(10.5)
        ph.font.bold = True
        ph.font.color.rgb = RGBColor(13, 148, 136)
        pd = tf.add_paragraph()
        pd.text = f"   {d}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 7: Sequence Diagram - Request Routing
    # ---------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "System Design: End-to-End Sequence Diagram")

    slide7.shapes.add_picture(diagram_paths["sequence_diagram"], Inches(0.8), Inches(1.4), width=Inches(7.8))

    add_card(slide7, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.5))
    tbox = slide7.shapes.add_textbox(Inches(9.0), Inches(1.6), Inches(3.3), Inches(5.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Sequence Execution Steps"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    seq_steps = [
        "1. Ingestion: Client submits query marked with surface provenance.",
        "2. Scoring: Classifiers generate multi-dimensional risk scores.",
        "3. Policy Check: DecisionEngine matches scores against policy rules.",
        "4. Conditional Attribution: If BLOCK or REVIEW, beam search finds minimal flip tokens.",
        "5. Sanitization: System excises offending words into a clean variant.",
        "6. Audit & Response: All metrics logged to SQLite before returning result."
    ]
    for s in seq_steps:
        ps = tf.add_paragraph()
        ps.text = s
        ps.font.size = Pt(9.5)
        ps.font.color.rgb = RGBColor(51, 65, 85)

    # ---------------------------------------------------------
    # SLIDE 8: Sequence Diagram - Runtime Policy Hot-Reloading
    # ---------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "System Design: Dynamic Policy Hot-Reloading Flow")

    add_card(slide8, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tbox_l = slide8.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(4.9))
    tfl = tbox_l.text_frame
    pl = tfl.paragraphs[0]
    pl.text = "The Zero-Retraining Paradigm"
    pl.font.size = Pt(16)
    pl.font.bold = True
    pl.font.color.rgb = RGBColor(30, 58, 138)

    points_l = [
        ("Base Paper Limitation (Llama Guard)", "Llama Guard fixes decision boundaries inside model weights. Changing a safety rule or adjusting sensitivity requires fine-tuning and redeployment."),
        ("CipherGuard Innovation", "Decouples classifier probabilities from organizational thresholds stored in an external JSON policy matrix."),
        ("Hot-Reload Mechanism", "PolicyLoader validates updates via Pydantic and applies threshold changes in memory immediately."),
        ("Zero Downtime", "Allows live security teams to tighten or loosen moderation policies on the fly without rebooting services.")
    ]
    for h, d in points_l:
        p_h = tfl.add_paragraph()
        p_h.text = f"• {h}"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(37, 99, 235)
        p_d = tfl.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = RGBColor(71, 85, 105)

    add_card(slide8, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tbox_r = slide8.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.9))
    tfr = tbox_r.text_frame
    pr = tfr.paragraphs[0]
    pr.text = "Hot-Reload Interaction Flow"
    pr.font.size = Pt(16)
    pr.font.bold = True
    pr.font.color.rgb = RGBColor(13, 148, 136)

    flow_steps = [
        ("Step 1: Admin PUT /policy", "Administrator sends updated JSON payload specifying new threshold (e.g., Jailbreak threshold 0.45 -> 0.35)."),
        ("Step 2: Pydantic Validation", "FastAPI validates the schema against PolicyConfig schema; rejects malformed requests immediately."),
        ("Step 3: Atomic Disk Write", "PolicyLoader commits new JSON configuration to disk atomically."),
        ("Step 4: Cache Update", "Internal cached_config and last_mtime timestamps update synchronously."),
        ("Step 5: Immediate Effect", "The next request entering route() evaluates against the newly loaded threshold with zero restart.")
    ]
    for h, d in flow_steps:
        p_h = tfr.add_paragraph()
        p_h.text = f"• {h}"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(15, 23, 42)
        p_d = tfr.add_paragraph()
        p_d.text = f"   {d}"
        p_d.font.size = Pt(9.8)
        p_d.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 9: State Chart Diagram
    # ---------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "System Design: Prompt Lifecycle State Chart Diagram")

    slide9.shapes.add_picture(diagram_paths["statechart_diagram"], Inches(0.8), Inches(1.4), width=Inches(7.8))

    add_card(slide9, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.5))
    tbox = slide9.shapes.add_textbox(Inches(9.0), Inches(1.6), Inches(3.3), Inches(5.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "State Lifecycle Description"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    state_desc = [
        ("Ingestion State", "Tags surface provenance (direct vs indirect) and normalizes input text."),
        ("Classification State", "Computes multi-dimensional risk scores via ensemble prediction."),
        ("Policy Evaluation State", "Compares risk vector against active policy thresholds."),
        ("Allow State", "Clears benign prompts directly to downstream LLM agent."),
        ("Flagged State", "Conducts beam search over token removals, computes delta score, and sanitizes input."),
        ("Error State", "Fails closed to safety block if unhandled exceptions occur.")
    ]
    for h, d in state_desc:
        ph = tf.add_paragraph()
        ph.text = f"• {h}:"
        ph.font.size = Pt(10.5)
        ph.font.bold = True
        ph.font.color.rgb = RGBColor(217, 119, 6)
        pd = tf.add_paragraph()
        pd.text = f"   {d}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 10: Deployment Diagram
    # ---------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "System Design: Physical Deployment Diagram")

    slide10.shapes.add_picture(diagram_paths["deployment_diagram"], Inches(0.8), Inches(1.4), width=Inches(7.8))

    add_card(slide10, Inches(8.8), Inches(1.4), Inches(3.7), Inches(5.5))
    tbox = slide10.shapes.add_textbox(Inches(9.0), Inches(1.6), Inches(3.3), Inches(5.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Deployment Architecture"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)

    dep_desc = [
        ("Presentation Layer", "Streamlit UI on Port 8501 for interactive testing and policy tuning."),
        ("API Gateway Layer", "FastAPI / Uvicorn on Port 8000 handling high-throughput REST calls."),
        ("Inference Tier", "GPU/CPU workers running Model Family A, Family B, and DeBERTa-v3."),
        ("Persistence Tier", "policy.json for live rules; SQLite with WAL mode for audit logging."),
        ("Reverse Proxy", "Nginx handles TLS termination and traffic routing.")
    ]
    for h, d in dep_desc:
        ph = tf.add_paragraph()
        ph.text = f"• {h}:"
        ph.font.size = Pt(10.5)
        ph.font.bold = True
        ph.font.color.rgb = RGBColor(22, 163, 74)
        pd = tf.add_paragraph()
        pd.text = f"   {d}"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = RGBColor(71, 85, 105)

    # ---------------------------------------------------------
    # SLIDE 11: Traceability Matrix (SRS vs Design vs Codebase)
    # ---------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "Implementation Verification: Traceability Matrix")

    rows = [
        ("FR-1: Dual-Surface Ingestion", "Ingestion Layer (DirectInput / IndirectInput)", "src/ingestion/", "Verified (tests/test_router.py)"),
        ("FR-2: Multi-Model Scoring", "ModelFamilyA, ModelFamilyB, DeBERTa-v3", "src/classification/", "Verified (tests/test_classification.py)"),
        ("FR-3: Dynamic Policy Matrix", "PolicyLoader, DecisionEngine (Hot-Reload)", "src/policy/, policy/policy.json", "Verified (tests/test_policy_engine.py)"),
        ("FR-4: Removal Token Attribution", "run_removal_beam_search, ContrastiveExplanation", "src/attribution/", "Verified (tests/test_attribution.py)"),
        ("FR-5: Audit Store & REST API", "AuditLogger (SQLite WAL), FastAPI (/route)", "src/logging_store.py, api/main.py", "Verified (FastAPI /route & /logs)"),
    ]

    # Draw Table
    x_start = Inches(0.8)
    y_start = Inches(1.5)
    col_widths = [Inches(3.2), Inches(3.6), Inches(2.7), Inches(2.2)]
    
    # Header Row
    hx = x_start
    headers = ["SRS Requirement", "System Design Component", "Codebase Module", "Verification Status"]
    for i, h in enumerate(headers):
        w = col_widths[i]
        card = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, hx, y_start, w, Inches(0.55))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.fill.background()
        
        tb = slide11.shapes.add_textbox(hx, y_start + Inches(0.08), w, Inches(0.4))
        p = tb.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        hx += w

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        ry = y_start + Inches(0.55 + r_idx * 0.95)
        rx = x_start
        bg_col = RGBColor(255, 255, 255) if r_idx % 2 == 0 else RGBColor(241, 245, 249)
        
        for c_idx, cell_text in enumerate(row_data):
            cw = col_widths[c_idx]
            card = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, rx, ry, cw, Inches(0.95))
            card.fill.solid()
            card.fill.fore_color.rgb = bg_col
            card.line.color.rgb = RGBColor(226, 232, 240)
            card.line.width = Pt(1)

            tb = slide11.shapes.add_textbox(rx + Inches(0.1), ry + Inches(0.1), cw - Inches(0.2), Inches(0.75))
            tb.text_frame.word_wrap = True
            p = tb.text_frame.paragraphs[0]
            p.text = cell_text
            p.font.size = Pt(10)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = RGBColor(30, 58, 138)
            elif c_idx == 3:
                p.font.bold = True
                p.font.color.rgb = RGBColor(22, 163, 74)
            else:
                p.font.color.rgb = RGBColor(51, 65, 85)
            rx += cw

    # ---------------------------------------------------------
    # SLIDE 12: Project Review Summary & Next Steps
    # ---------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Summary & Next Steps for Project Review 2")

    add_card(slide12, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    tbox_l = slide12.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(4.9))
    tfl = tbox_l.text_frame
    pl = tfl.paragraphs[0]
    pl.text = "Deliverables Completed"
    pl.font.size = Pt(16)
    pl.font.bold = True
    pl.font.color.rgb = RGBColor(30, 58, 138)

    deliverables = [
        "1. Software Requirements Specification (SRS) written following IEEE 830 standards.",
        "2. System Design Architecture detailing the 5 decoupled sub-pipelines.",
        "3. Complete UML Diagrams: Class, Sequence (Routing & Hot-Reload), State Chart, and Deployment diagrams.",
        "4. Full test suite passing with 100% success (13/13 unit & integration tests).",
        "5. Interactive demonstration environment ready via Streamlit (:8501) and FastAPI (:8000)."
    ]
    for d in deliverables:
        pd = tfl.add_paragraph()
        pd.text = f"• {d}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = RGBColor(51, 65, 85)

    add_card(slide12, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    tbox_r = slide12.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.9))
    tfr = tbox_r.text_frame
    pr = tfr.paragraphs[0]
    pr.text = "Live Demo & Future Milestones"
    pr.font.size = Pt(16)
    pr.font.bold = True
    pr.font.color.rgb = RGBColor(13, 148, 136)

    milestones = [
        "• Live Demonstration:",
        "   1. Submit a direct jailbreak attempt -> inspect BLOCK & minimal tokens removed.",
        "   2. Update threshold dynamically via PUT /policy -> verify instant policy shift.",
        "   3. Ingest indirect document attack -> observe decoupled surface calibration.",
        "• Phase 3 Planned Enhancements:",
        "   1. Large-scale benchmark evaluation (InjecAgent & AgentDojo).",
        "   2. Batched beam search GPU acceleration to push attribution under 100ms.",
        "   3. Multi-turn conversation state tracking and output moderation."
    ]
    for m in milestones:
        pm = tfr.add_paragraph()
        pm.text = m
        pm.font.size = Pt(10.5)
        pm.font.color.rgb = RGBColor(15, 23, 42) if m.startswith("•") else RGBColor(71, 85, 105)
        if m.startswith("•"):
            pm.font.bold = True

    # Save presentation
    output_pptx = "CipherGuard_Project_Review_2.pptx"
    prs.save(output_pptx)
    print(f"Presentation successfully saved to: {output_pptx}")
    return output_pptx


if __name__ == "__main__":
    print("Generating diagram figures...")
    paths = {
        "system_design": create_system_design_diagram(),
        "class_diagram": create_class_diagram(),
        "sequence_diagram": create_sequence_diagram(),
        "statechart_diagram": create_statechart_diagram(),
        "deployment_diagram": create_deployment_diagram(),
    }
    print("Generating PowerPoint deck...")
    deck_path = build_presentation(paths)
    print(f"Done! Deck available at: {deck_path}")
