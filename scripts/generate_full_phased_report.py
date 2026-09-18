"""
Script to generate the comprehensive CipherGuard Project Status Report DOCX document,
incorporating all 4 phases from the phased prompt specification:
- Phase 1: Dataset Engineering (Completed)
- Phase 2: Models, Ensemble, Thresholds (Upcoming)
- Phase 3: Calibration, Policy/Security Tests, Attribution, Ablations (Upcoming)
- Phase 4: Reproducibility, Regression Testing, Documentation, Final Report (Upcoming)

Strictly omits all personal, team member, guide, and institutional names.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

HEX_PRIMARY = "0F2942"      # Deep Navy
HEX_SECONDARY = "1E5F8A"    # Steel Blue
HEX_ACCENT = "0D9488"       # Teal
HEX_DARK = "1E293B"         # Charcoal Text
HEX_LIGHT_BG = "F1F5F9"     # Light Slate Background
HEX_BORDER = "CBD5E1"       # Border Grey
HEX_AMBER_BG = "FEF3C7"     # Warning / Amber
HEX_GREEN_BG = "DCFCE7"     # Success / Green

COLOR_PRIMARY = RGBColor(15, 41, 66)
COLOR_SECONDARY = RGBColor(30, 95, 138)
COLOR_DARK = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_full_phased_docx(filename):
    doc = Document()
    
    # 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("CipherGuard LLM Security Layer | Project Status Report (Phases 1-4)")
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = COLOR_MUTED

    # Title & Subtitle
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("CipherGuard: Comprehensive Project Status & Phased Roadmap Report")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Detailed Technical Audit: Inception, Phase 1 Completion, Identified Constraints, and Phased Roadmap (Phases 2-4)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata Strip Box
    meta_table = doc.add_table(rows=1, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    col_widths = [Inches(1.75), Inches(1.75), Inches(1.75), Inches(1.75)]
    meta_data = [
        ("Current Phase", "Phase 1 Completed\n(Dataset Engineering)"),
        ("Live Endpoints", "FastAPI REST API &\nStreamlit UI Active"),
        ("Verification Suite", "13 / 13 Tests Passed\n(100% Pass Rate)"),
        ("Phased Scope", "Phases 1-4 Detailed\n(Zero Retraining Router)")
    ]
    for i, (k, v) in enumerate(meta_data):
        cell = meta_table.rows[0].cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, HEX_LIGHT_BG)
        set_cell_margins(cell, 100, 100, 120, 120)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        r_k = cp.add_run(k.upper() + "\n")
        r_k.font.size = Pt(8)
        r_k.font.bold = True
        r_k.font.color.rgb = COLOR_MUTED
        r_v = cp.add_run(v)
        r_v.font.size = Pt(8.5)
        r_v.font.bold = True
        r_v.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper formatters
    def add_section_header(text, level=1):
        p = doc.add_paragraph()
        if level == 1:
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY
        elif level == 3:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK

    def add_body(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_DARK

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------------------
    # 1. Executive Summary
    # -------------------------------------------------------------------------
    add_section_header("1. Executive Summary & Core Motivation", 1)
    add_body(
        "CipherGuard is engineered as an auditable, pre-agent security gateway positioned inline between untrusted "
        "inputs and core generative Large Language Model (LLM) agents. Enterprise LLM applications are vulnerable to "
        "two distinct attack surfaces: Direct Prompt Injection (DPI / jailbreaks in user turns) and Indirect Prompt "
        "Injection (IPI delivered via third-party retrieved context, tool outputs, emails, database records, and web data). "
        "Existing defensive mechanisms exhibit critical operational deficiencies:"
    )
    add_bullet(
        "Static safety classifiers (such as Llama Guard or ShieldGemma) hardcode moderation criteria directly "
        "into model parameters. Adjusting a threshold or responding to evolving risk tolerance requires expensive, "
        "time-consuming model retraining.",
        bold_prefix="Rigidity of Monolithic Classifiers: "
    )
    add_bullet(
        "Using secondary LLMs to judge input safety introduces high per-token API costs, latency overheads "
        "of several hundred milliseconds to multiple seconds, and non-deterministic, unauditable rejection rationales.",
        bold_prefix="Latency & Cost of LLM-as-a-Judge: "
    )
    add_body(
        "CipherGuard resolves these bottlenecks through three interconnected architectural pillars:"
    )
    add_bullet(
        "A JSON-configured decision engine that maps risk thresholds and routing actions (BLOCK, REVIEW, ALLOW) "
        "independently of classifier weights, enabling zero-retraining runtime policy adjustments.",
        bold_prefix="1. Dynamic Policy Matrix: "
    )
    add_bullet(
        "Independent risk scoring and threshold calibration tailored to the distinct statistical distributions "
        "of direct user turns versus indirect retrieved context.",
        bold_prefix="2. Dual-Surface Threat Modeling: "
    )
    add_bullet(
        "An approximate beam search counterfactual engine that isolates the minimal token subset whose "
        "removal flips a BLOCK decision to PASS, delivering an actionable audit rationale and sanitized safe text.",
        bold_prefix="3. Removal-Based Contrastive Attribution: "
    )

    # -------------------------------------------------------------------------
    # 2. System Architecture & Subsystems
    # -------------------------------------------------------------------------
    add_section_header("2. System Architecture & Subsystems Decomposition", 1)
    add_body(
        "CipherGuard is partitioned into five isolated, modular subsystems operating in a coordinated pipeline:"
    )
    add_bullet(
        "DirectInput (user turns) and IndirectInput (RAG chunks, tool outputs, emails, database rows, web snippets) "
        "with end-to-end metadata and surface provenance preservation.",
        bold_prefix="1. Ingestion Layer (src/ingestion/): "
    )
    add_bullet(
        "Multi-model ensemble combining Model Family A (lexical TF-IDF + Logistic Regression), Model Family B "
        "(neural TF-IDF + 32-unit MLP), ProtectAI DeBERTa-v3 semantic injection classifier, ToxicBERT (Hate/Toxicity), "
        "and Microsoft Presidio Analyzer (PII Leakage). Scores are fused into a 5-dimensional risk vector via RiskAggregator.",
        bold_prefix="2. Classification & Risk Fusion (src/classification/): "
    )
    add_bullet(
        "A declarative rules engine evaluating policy/policy.json with Pydantic validation, wildcard fallback "
        "(surface-specific -> '*'), and atomic zero-downtime hot reloading.",
        bold_prefix="3. Dynamic Policy Decision Engine (src/policy/): "
    )
    add_bullet(
        "Executes upon BLOCK or REVIEW decisions using batched beam search (k=5, L=6) to identify minimal trigger "
        "tokens, score reduction delta (Delta s), and safe sanitized text.",
        bold_prefix="4. Token Attribution Engine (src/attribution/): "
    )
    add_bullet(
        "SQLite WAL persistence (cipherguard_audit.db) logging text, risk vectors, rules, latency, and attribution; "
        "FastAPI REST server (:8000); and Streamlit dashboard (:8501).",
        bold_prefix="5. Persistence & Interface Layer: "
    )

    # -------------------------------------------------------------------------
    # 3. Status of Inception & Proof-of-Concept Baseline
    # -------------------------------------------------------------------------
    add_section_header("3. Status of Inception & Proof-of-Concept Baseline", 1)
    add_body(
        "The project began with an initial pilot / proof-of-concept dataset (n=85: 45 safe, 40 unsafe; 55 direct, "
        "30 indirect). Key experimental findings from the initial baseline established the foundation:"
    )
    add_bullet(
        "Repeated 5x5 stratified cross-validation confirmed strong separability on the synthetic pilot. "
        "Model Family A achieved 0.962 accuracy and 0.959 precision; Model Family B achieved 0.899 accuracy and 0.970 recall, "
        "demonstrating that model family choice inherently trades precision for recall.",
        bold_prefix="Model Families Trade-off (Table I Reproduction): "
    )
    add_bullet(
        "Under a narrow budget (b=5, k<=6), linear models achieved 0.0% flips. Under a wide budget (b=10, k<=10), "
        "12.5% flipped but required an average of 9.2 tokens removed, demonstrating diffused lexical importance in bag-of-words models.",
        bold_prefix="Attribution Budget Ablation (Table I Reproduction): "
    )
    add_bullet(
        "Zero-shot ProtectAI DeBERTa-v3 achieved 0.977 accuracy, 100% precision, and confirmed surface asymmetry "
        "(92% recall on direct vs 100% on indirect). Transformer beam search achieved an 84.2% flip rate with an average "
        "of only 2.16 tokens removed (highly compact, actionable counterfactuals).",
        bold_prefix="Contextual Transformer Benchmark (Table II Reproduction): "
    )

    # -------------------------------------------------------------------------
    # 4. Phase 1 Accomplishments: Dataset Engineering (COMPLETED)
    # -------------------------------------------------------------------------
    add_section_header("4. Phase 1 Accomplishments: Dataset Engineering (COMPLETED)", 1)
    add_body(
        "Phase 1 focused on replacing the initial 85-sample pilot with a multi-label dataset pipeline, resolving "
        "prevalence gaps, eliminating data leakage, and establishing strict data governance. All Phase 1 tasks are complete:"
    )
    
    add_section_header("Section 1: Preservation of the n=85 Pilot Dataset", 2)
    add_bullet(
        "pilot_dataset_85.json is locked via SHA-256 checksum (89d6247b9a92ae89...) in .pilot_checksum. "
        "It is treated strictly as 'Pilot / proof-of-concept dataset (n=85)' and is never used as the training or final evaluation split.",
        bold_prefix="Pilot Dataset Lock: "
    )

    add_section_header("Sections 2 & 5: Expanded Multi-Label Dataset (N = 923 Samples)", 2)
    add_body(
        "The expanded dataset aggregates six sources into a unified multi-label schema with five non-mutually exclusive categories:"
    )
    
    # Table of dataset composition
    t_ds = doc.add_table(rows=8, cols=3)
    t_ds.alignment = WD_TABLE_ALIGNMENT.CENTER
    ds_widths = [Inches(3.2), Inches(1.1), Inches(2.7)]
    ds_headers = ["Source Repository / Synthetic Generator", "Count", "Role in Multi-Label Pipeline"]
    ds_rows = [
        ["deepset/prompt-injections (Hugging Face)", "350", "Public direct prompt injection benchmark"],
        ["tweet_eval/hate (Hugging Face)", "200", "Public hate and toxicity benchmark (100 hate, 100 benign)"],
        ["cipherguard_maltools_synthetic", "125", "Multi-turn malicious tool attacks & benign controls"],
        ["pilot_dataset_85.json (PoC)", "85", "Original pilot converted to multi-label format (read-only)"],
        ["cipherguard_pii_synthetic", "99", "4-tier synthetic PII leakage dataset"],
        ["cipherguard_indirect_synthetic", "64", "Indirect surface attacks (RAG, email, DB, web, tools)"],
        ["Total Combined Dataset", "923", "Full Expanded Multi-Label Corpus"]
    ]
    for j, h in enumerate(ds_headers):
        c = t_ds.rows[0].cells[j]
        c.width = ds_widths[j]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(ds_rows):
        for j, val in enumerate(row):
            c = t_ds.rows[i+1].cells[j]
            c.width = ds_widths[j]
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            if i == len(ds_rows) - 1:
                bg = "E2E8F0"
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            r.font.bold = (i == len(ds_rows) - 1)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_section_header("Section 3: Resolution of the PII Dataset Gap", 2)
    add_bullet(
        "The legacy 829-sample dataset contained zero PII positives, leaving the category untrainable. Phase 1 created 68 positive "
        "samples across a 4-tier taxonomy (Tier 1: Benign PII discussion, Tier 2: Extraction requests, Tier 3: In-text sensitive data, "
        "Tier 4: Malicious credential harvesting).",
        bold_prefix="Synthetic 4-Tier PII Taxonomy: "
    )
    add_bullet(
        "All values follow demonstrably fake standards: SSN area code 000 (invalid per IRS/SSA rules), Luhn-failing credit card numbers, "
        "_FAKE_ API key prefixes, 555-xxxx phone numbers, and RFC 2606 .invalid email TLDs.",
        bold_prefix="Demonstrably Non-Real Entities: "
    )

    add_section_header("Section 4: Improved Malicious Tools Dataset", 2)
    add_bullet(
        "Expanded from trivial keyword memorization (e.g. 'rm -rf') to contextual attack patterns including shell commands, "
        "SQL injection patterns, reverse-shell payloads, tool output abuse, and credential extraction, paired with benign system admin commands.",
        bold_prefix="Contextual Tool Invocations: "
    )

    add_section_header("Section 5: Surface Metadata & Distribution Alignment", 2)
    add_bullet(
        "100% of samples are tagged with attack_type and severity (low: 535, medium: 122, high: 130, critical: 136). "
        "Indirect surface share was expanded from 3.6% (30 samples) to 10.2% (94 samples).",
        bold_prefix="Surface Representation: "
    )

    add_section_header("Section 6 & 19: Deterministic Stratified Split & Zero-Leakage Audit", 2)
    add_bullet(
        "data/train.json: 645 samples (69.9%), data/val.json: 138 samples (15.0%), data/test.json: 140 samples (15.2%). "
        "Stratification ensures label prevalence variance is strictly <= 0.7% across splits (split_manifest.json).",
        bold_prefix="70 / 15 / 15 Partition (Seed 42): "
    )
    add_bullet(
        "Pairwise 3-gram character Jaccard similarity audit over ~198,000 comparisons revealed zero malicious or positive-label near-duplicate leakage across splits (leakage_report.json).",
        bold_prefix="Data Leakage Audit: "
    )
    add_bullet(
        "data/test.json is strictly held out and reserved exclusively for final evaluation reporting in Phase 4.",
        bold_prefix="Held-Out Test Integrity: "
    )

    add_section_header("Test Suite Verification Status", 2)
    add_body("The complete test suite was executed in the project virtual environment with 100% pass rate:")
    add_bullet("13 / 13 tests passed in 57.38s (test_attribution.py: 2, test_classification.py: 4, test_policy_engine.py: 3, test_router.py: 4).")

    # -------------------------------------------------------------------------
    # 5. Identified Limitations, Bottlenecks & Pendings
    # -------------------------------------------------------------------------
    add_section_header("5. Identified Limitations, Bottlenecks & Pendings", 1)
    add_body(
        "In alignment with strict research honesty, the following limitations and technical pendings are documented:"
    )
    
    t_lim = doc.add_table(rows=6, cols=3)
    t_lim.alignment = WD_TABLE_ALIGNMENT.CENTER
    lim_widths = [Inches(1.8), Inches(2.6), Inches(2.6)]
    lim_headers = ["Identified Limitation", "Root Cause & Operational Impact", "Planned Mitigation / Next Action"]
    lim_rows = [
        [
            "CPU Attribution Latency (4-6s)",
            "Evaluating candidate token deletions iteratively through DeBERTa-v3 on CPU creates latency bottlenecks during explanations.",
            "Deploy GPU acceleration, ONNX Runtime INT8 quantization, or fast two-stage token pre-filtering."
        ],
        [
            "Model Retraining on Train Split",
            "Existing pickle models (models/family_*.pkl) were fitted prior to the strict 70/15/15 split on all 923 samples.",
            "Execute evaluation/phase2_train.py to retrain Family A and B strictly on data/train.json (645 samples)."
        ],
        [
            "Archived LoRA DeBERTa Model",
            "Archived LoRA fine-tuned DeBERTa (~704 samples) produced false positives. Pretrained ProtectAI checkpoint retained.",
            "Keep archived; do not claim it is superior. Revisit fine-tuning only when larger curated corpora are available."
        ],
        [
            "Diffused Signal in Linear Models",
            "TF-IDF models diffuse signal across words, requiring ~9.2 token removals vs 2.16 in DeBERTa.",
            "Use transformer backends when concise, human-interpretable explanations are required."
        ],
        [
            "PowerShell Environment Execution",
            "Running 'pytest' directly without activating the virtual environment triggers command not found.",
            "Standardize on .\\.venv\\Scripts\\pytest.exe or activate virtualenv first."
        ]
    ]
    for j, h in enumerate(lim_headers):
        c = t_lim.rows[0].cells[j]
        c.width = lim_widths[j]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(lim_rows):
        for j, val in enumerate(row):
            c = t_lim.rows[i+1].cells[j]
            c.width = lim_widths[j]
            bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------------------
    # 6. Detailed Roadmap: Phase 2 — Models, Ensemble, Thresholds
    # -------------------------------------------------------------------------
    add_section_header("6. Detailed Roadmap: Phase 2 — Models, Ensemble, Thresholds", 1)
    add_body(
        "Phase 2 builds upon the verified Phase 1 dataset partitions to train discriminators strictly on train.json, "
        "benchmark individual models on val.json, ablate ensemble fusion strategies, and calibrate risk thresholds. "
        "Ground rule: data/test.json is strictly held out and never inspected during Phase 2."
    )

    add_section_header("Section 7: True Held-Out Test Set Enforcement", 2)
    add_bullet(
        "Confirm that data/test.json (140 samples) remains completely isolated from training, hyperparameter tuning, "
        "threshold search, and ensemble weight learning.",
        bold_prefix="Quarantine Protocol: "
    )

    add_section_header("Section 8: Evaluate All Individual Models on Validation Split", 2)
    add_body("Benchmark all candidate classifiers on data/val.json (138 samples) across applicable categories:")
    add_bullet("1. Model Family A (TF-IDF + Logistic Regression, retrained on train.json).")
    add_bullet("2. Model Family B (TF-IDF + MLP, retrained on train.json).")
    add_bullet("3. ProtectAI DeBERTa-v3 (zero-shot, Jailbreak and Prompt Injection).")
    add_bullet("4. ToxicBERT (unitary/toxic-bert, Hate/Toxicity only).")
    add_bullet("5. Presidio Analyzer (Microsoft Presidio, PII Leakage only).")
    add_bullet("6. Current Ensemble baseline (MAX fusion across models).")
    add_body(
        "Metrics to compute per model per category: Accuracy, Precision, Recall, F1-Score, Confusion Matrix (TP, FP, TN, FN), "
        "ROC-AUC, and PR-AUC. Security classification prioritizes Recall and Precision/Recall trade-offs over raw accuracy."
    )

    add_section_header("Section 9: Surface-Decoupled Evaluation (Direct vs. Indirect)", 2)
    add_bullet(
        "Separately calculate Precision, Recall, and F1 for DIRECT and INDIRECT surfaces on Jailbreak and Prompt Injection. "
        "Do not report a single monolithic number; empirical surface asymmetry is central to CipherGuard's architectural thesis.",
        bold_prefix="Surface Disaggregation: "
    )

    add_section_header("Section 10: Scientific Protocol on DeBERTa Fine-Tuning", 2)
    add_bullet(
        "The archived LoRA fine-tuned DeBERTa (~704 samples) produced unacceptable false positives. Maintain it as archived; "
        "do not claim the in-house model is superior. Document why the pretrained ProtectAI checkpoint is preferred and what dataset "
        "scale/quality would justify revisiting fine-tuning.",
        bold_prefix="Honest Model Governance: "
    )

    add_section_header("Section 11: Experimental Ensemble Fusion Ablation", 2)
    add_body(
        "The current baseline uses conservative MAX fusion (Final Score = max(Model A, Model B, Specialist)). "
        "Phase 2 will systematically execute an ablation study across seven candidate fusion strategies:"
    )
    add_bullet("Strategy A: Logistic Regression only (Model Family A baseline).")
    add_bullet("Strategy B: MLP only (Model Family B baseline).")
    add_bullet("Strategy C: DeBERTa / Specialist models only.")
    add_bullet("Strategy D: LR + MLP blend.")
    add_bullet("Strategy E: LR + MLP + Specialist with MAX fusion (current operational baseline).")
    add_bullet("Strategy F: LR + MLP + Specialist with Weighted Average fusion (weights tuned on validation predictions).")
    add_bullet("Strategy G: Small meta-classifier / Stacking model (fitted on validation score vectors).")
    add_body(
        "All weights and meta-models are learned exclusively using validation predictions. The winning strategy is selected "
        "based on validation F1/Recall, and only then evaluated on the held-out test split."
    )

    add_section_header("Section 12: Risk Score Calibration & Threshold Search", 2)
    add_bullet(
        "Current static thresholds in policy.json (BLOCK >= 0.85, REVIEW 0.45-0.84, ALLOW < 0.45 or category-specific values) "
        "must not be assumed optimal. Execute a systematic grid search on validation data.",
        bold_prefix="Validation Threshold Search: "
    )
    add_bullet(
        "Objective function: Minimize dangerous false negatives (missed attacks) while maintaining false positives below an acceptable ceiling. "
        "Determine separate thresholds for direct vs indirect surfaces where supported by validation data. Freeze selected thresholds before touching the test split.",
        bold_prefix="Optimization Criterion: "
    )

    # -------------------------------------------------------------------------
    # 7. Detailed Roadmap: Phase 3 — Calibration, Policy Tests, Attribution, Ablations
    # -------------------------------------------------------------------------
    add_section_header("7. Detailed Roadmap: Phase 3 — Calibration, Security Tests, Attribution & Ablations", 1)
    add_body(
        "Phase 3 focuses on probability calibration, formal security scenario unit testing, explainability benchmarking, "
        "comprehensive ablation studies, error analysis, and external benchmark compatibility."
    )

    add_section_header("Section 13: Probability Calibration Analysis", 2)
    add_bullet(
        "Evaluate whether model output scores represent true risk probabilities using reliability diagrams, calibration curves, "
        "Brier score, and Expected Calibration Error (ECE). Experiment with Platt scaling (logistic calibration) or isotonic regression "
        "fitted on validation data.",
        bold_prefix="ECE & Brier Score: "
    )

    add_section_header("Section 14: Automated Policy Matrix Unit Testing", 2)
    add_bullet(
        "Validate the strict separation of model risk scoring ('what is the risk?') from policy resolution ('what action?').",
        bold_prefix="Model/Policy Separation: "
    )
    add_bullet(
        "Develop unit tests covering low/medium/high risk -> correct action mapping, threshold boundary values, surface-specific matching, "
        "conflicting multi-category risk resolutions, and verify that updating policy.json instantly modifies routing behavior without model retraining.",
        bold_prefix="Policy Unit Tests: "
    )

    add_section_header("Section 15: Dual-Surface Security Scenario Test Fixtures", 2)
    add_bullet(
        "Construct controlled evaluation fixtures simulating actual indirect attack pathways: (1) RAG retrieved document chunk injection, "
        "(2) Malicious webpage content injection, (3) Malicious email body injection, (4) Injected database record retrieval, and "
        "(5) Malicious tool-output payload.",
        bold_prefix="Five Injection Fixtures: "
    )
    add_bullet(
        "Demonstrate that attacks embedded in third-party context are successfully intercepted before downstream LLM/Agent invocation.",
        bold_prefix="Pre-Agent Verification: "
    )

    add_section_header("Section 16: Removal-Based Attribution Benchmarking", 2)
    add_bullet(
        "Systematically evaluate beam search counterfactuals across blocked instances. Measure: decision flip rate, mean tokens removed, "
        "search execution latency, success rate by attack type, and success rate by surface.",
        bold_prefix="Attribution Metrics: "
    )
    add_bullet(
        "Ablate beam search budgets: Narrow (beam=5, L=6) vs. Wide (beam=10, L=10). Explicitly describe output in all reporting as an "
        "approximate contrastive explanation, not a provably exact minimal set.",
        bold_prefix="Budget Ablation & Framing: "
    )

    add_section_header("Section 17: Comprehensive Multi-Dimensional Ablation Matrix", 2)
    add_bullet("1. Dataset Ablation: Pilot only vs Public only vs Public+Custom vs Public+Custom+Synthetic.")
    add_bullet("2. Model Ablation: Family A (LR) vs Family B (MLP) vs DeBERTa vs Specialists vs Ensembles.")
    add_bullet("3. Fusion Ablation: MAX vs Weighted Average vs Stacking meta-classifier.")
    add_bullet("4. Policy Ablation: Monolithic shared threshold vs Decoupled surface-specific thresholds.")
    add_bullet("5. Attribution Ablation: Candidate beam sizes (k=3, 5, 10) and removal caps (L=4, 6, 10).")

    add_section_header("Section 18: Confusion Matrix & False Positive / Negative Error Analysis", 2)
    add_bullet(
        "Generate and save confusion matrices for all five safety categories. Perform qualitative inspection of representative "
        "false-positive and false-negative instances, logging input text, ground-truth label, predicted vector, surface, attack type, "
        "and router decision to identify systemic model failure modes.",
        bold_prefix="Qualitative Error Audit: "
    )

    add_section_header("Section 20: External / Held-Out Benchmark Support", 2)
    add_bullet(
        "Add evaluation-only adapters for recognized academic benchmarks where accessible: WildGuardTest, JailbreakBench, HarmBench, "
        "and InjecAgent/AgentDojo. External benchmarks must be strictly eval-only and never merged into training. "
        "If a benchmark cannot be downloaded in this environment, build the harness code and document the constraint explicitly.",
        bold_prefix="Standardized Academic Benchmarks: "
    )

    # -------------------------------------------------------------------------
    # 8. Detailed Roadmap: Phase 4 — Reproducibility, Documentation & Final Deliverables
    # -------------------------------------------------------------------------
    add_section_header("8. Detailed Roadmap: Phase 4 — Reproducibility, Documentation & Final Deliverables", 1)
    add_body(
        "Phase 4 ensures enterprise-grade reproducibility, expands regression testing, updates technical documentation, "
        "verifies all 16 project deliverables, and generates the final academic implementation report."
    )

    add_section_header("Section 21: Reproducibility & Central Experiment Configuration", 2)
    add_bullet(
        "Establish a central configuration file (configs/experiment.yaml) specifying random seeds, dataset versions, train/val/test counts, "
        "hyperparameters, frozen thresholds, and fusion methods to prevent hard-coded discrepancies.",
        bold_prefix="Central Config: "
    )

    add_section_header("Section 22: Documentation Integrity Standards", 2)
    add_bullet(
        "Enforce strict clarity in all documentation: clearly distinguish between CURRENT IMPLEMENTATION, PILOT POC RESULTS (n=85), "
        "and FINAL EVALUATION. Never claim n=85 validates real-world performance. Document the PII dataset gap fix, the archived DeBERTa "
        "rationale, threshold calibration findings, and direct-vs-indirect results.",
        bold_prefix="Documentation Standards: "
    )

    add_section_header("Section 23: Application Functionality & Regression Testing", 2)
    add_bullet(
        "Ensure all core interfaces remain stable. Run existing test suite and introduce regression tests for: dataset schema validation, "
        "train/test leakage checks, PII entity detection, model score bounds, ensemble fusion, dynamic threshold routing, "
        "direct/indirect handling, token attribution, and FastAPI REST endpoint contracts.",
        bold_prefix="Regression Suite: "
    )

    add_section_header("Section 24: Engineering Constraints Verification", 2)
    add_bullet(
        "Confirm all systems operate under defined constraints: CPU-compatible baselines, no mandatory GPU dependency, "
        "lazy loading for DeBERTa/ToxicBERT, graceful degradation fallbacks, and reliable Windows 11 compatibility.",
        bold_prefix="Runtime Constraints: "
    )

    add_section_header("Section 25: Verification Checklist of 16 Final Deliverables", 2)
    add_body("Before final sign-off, verify the existence and empirical validity of each deliverable:")
    
    deliverables = [
        ("1. Updated Dataset Pipeline", "build_expanded_dataset.py combining 6 sources (N=923)"),
        ("2. Improved Dataset Files", "expanded_dataset.json, train.json, val.json, test.json"),
        ("3. Proper Train/Val/Test Split", "Stratified 70/15/15 split manifest (seed=42)"),
        ("4. PII-Positive Dataset", "4-tier synthetic PII leakage dataset (68 positives)"),
        ("5. Improved Malicious-Tool Dataset", "Contextual command attacks and benign controls"),
        ("6. Multi-Model Benchmark", "Evaluation metrics across Family A, B, DeBERTa, ToxicBERT, Presidio"),
        ("7. Ensemble Ablation Study", "Ablation of 7 fusion strategies (A-G) with selected winner"),
        ("8. Threshold Calibration", "Validation grid search optimizing per-surface operating points"),
        ("9. Probability Calibration Analysis", "ECE, Brier score, and reliability curves"),
        ("10. Direct vs. Indirect Evaluation", "Disaggregated metrics proving surface asymmetry"),
        ("11. Held-Out Test Evaluation", "Unbiased one-time evaluation on quarantined data/test.json"),
        ("12. Confusion Matrices & Error Analysis", "Per-category error breakdown and qualitative FP/FN analysis"),
        ("13. Attribution Evaluation", "Flip rate, tokens removed, search latency, and budget ablation"),
        ("14. Policy & Dual-Surface Security Tests", "Automated test fixtures for policy resolution and 5 injection surfaces"),
        ("15. External Benchmark Support", "Evaluation harness for WildGuardTest / InjecAgent"),
        ("16. Central Experiment Config & Docs", "configs/experiment.yaml, updated user manuals, and architecture specs")
    ]
    t_del = doc.add_table(rows=len(deliverables)+1, cols=2)
    t_del.alignment = WD_TABLE_ALIGNMENT.CENTER
    del_widths = [Inches(2.5), Inches(4.5)]
    for j, h in enumerate(["Required Deliverable", "Verification Standard & Artifact"]):
        c = t_del.rows[0].cells[j]
        c.width = del_widths[j]
        set_cell_background(c, HEX_SECONDARY)
        set_cell_margins(c, 60, 60, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for i, (item, std) in enumerate(deliverables):
        c0 = t_del.rows[i+1].cells[0]
        c1 = t_del.rows[i+1].cells[1]
        c0.width = del_widths[0]
        c1.width = del_widths[1]
        bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(item)
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(std)
        r1.font.size = Pt(8)
        r1.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_section_header("Section 26: Final Report Structure (Sections A through N)", 2)
    add_body(
        "The project will conclude with a formal academic implementation report structured into fourteen standardized sections: "
        "(A) What was changed, (B) Dataset statistics, (C) Train/validation/test statistics, (D) Per-category label distribution, "
        "(E) Model performance benchmarks, (F) Ensemble comparison & ablation, (G) Threshold optimization results, "
        "(H) Direct vs. indirect surface results, (I) Probability calibration results, (J) Attribution evaluation results, "
        "(K) External benchmark results, (L) False-positive and false-negative qualitative analysis, (M) Remaining limitations, "
        "and (N) Recommended next steps."
    )

    # -------------------------------------------------------------------------
    # 9. Teammate Quickstart & Local Instructions
    # -------------------------------------------------------------------------
    add_section_header("9. Local Quickstart & Instructions for Teammates", 1)
    add_body("To execute, test, and inspect the codebase locally, use the following standardized shell commands:")
    
    commands = [
        ("Activate Virtual Environment", ".\\.venv\\Scripts\\Activate.ps1"),
        ("Run Full Test Suite (13 tests)", ".\\.venv\\Scripts\\pytest.exe tests/ -v"),
        ("Launch Interactive Streamlit UI", ".\\.venv\\Scripts\\streamlit.exe run ui/app.py"),
        ("Start FastAPI REST Server", ".\\.venv\\Scripts\\uvicorn.exe api.main:app --reload --port 8000"),
        ("Run Phase 2 Model Retraining", ".\\.venv\\Scripts\\python.exe -m evaluation.phase2_train"),
        ("Run Phase 2 Validation Benchmark", ".\\.venv\\Scripts\\python.exe -m evaluation.phase2_benchmark"),
        ("Run Phase 2 Threshold Grid Search", ".\\.venv\\Scripts\\python.exe -m evaluation.phase2_threshold_search"),
        ("Run GPU DeBERTa LoRA Fine-Tuning", "python -m evaluation.finetune_deberta --mode lora  (on Colab/GPU)")
    ]
    t_cmd = doc.add_table(rows=len(commands)+1, cols=2)
    t_cmd.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_widths = [Inches(2.5), Inches(4.5)]
    for j, h in enumerate(["Target Operation", "Shell Command"]):
        c = t_cmd.rows[0].cells[j]
        c.width = cmd_widths[j]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, 60, 60, 80, 80)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for i, (op, cmd) in enumerate(commands):
        c0 = t_cmd.rows[i+1].cells[0]
        c1 = t_cmd.rows[i+1].cells[1]
        c0.width = cmd_widths[0]
        c1.width = cmd_widths[1]
        bg = HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_margins(c0, 50, 50, 80, 80)
        set_cell_margins(c1, 50, 50, 80, 80)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(op)
        r0.font.size = Pt(8)
        r0.font.color.rgb = COLOR_DARK
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(cmd)
        r1.font.name = "Consolas"
        r1.font.size = Pt(8)
        r1.font.color.rgb = COLOR_PRIMARY

    doc.save(filename)
    print(f"Full Phased DOCX successfully generated: {filename}")


if __name__ == "__main__":
    os.makedirs("docs", exist_ok=True)
    targets = [
        "CipherGuard_Phase1_Project_Status_Report.docx",
        "docs/CipherGuard_Phase1_Project_Status_Report.docx",
        "CipherGuard_Project_Status_Report_Phases_1_to_4.docx",
        "docs/CipherGuard_Project_Status_Report_Phases_1_to_4.docx"
    ]
    for target in targets:
        try:
            generate_full_phased_docx(target)
        except PermissionError:
            print(f"Notice: {target} is currently open in Word. Skipping locked file.")
        except Exception as e:
            print(f"Error saving {target}: {e}")

